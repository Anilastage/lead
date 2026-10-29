import csv
import hashlib
import io
import re
from pathlib import Path

import pandas as pd
from email_validator import EmailNotValidError, validate_email

ALIASES = {
    "name": [
        "name", "full_name", "business_name", "company", "company_name", "firm",
        "agency", "business", "organization", "organisation", "contact_name", "customer",
    ],
    "first_name": ["first_name", "fname", "given_name", "firstname"],
    "email": [
        "email", "e_mail", "email_id", "email_address", "mail", "emailaddress",
        "email_id_1", "work_email",
    ],
    "phone": [
        "phone", "phone_number", "mobile", "mobile_number", "whatsapp", "whatsapp_number",
        "contact", "contact_number", "number", "tel", "telephone", "cell", "cellphone",
        "mobile_no", "phone_no",
    ],
    "business_type": [
        "business_type", "business_category", "category", "industry", "niche", "sector",
        "type", "service", "services", "btype",
    ],
    "website": ["website", "website_url", "site", "url", "web", "domain"],
    "city": ["city", "location", "area", "address", "region", "locality"],
    "rating": ["rating", "google_rating", "stars", "score", "google_stars"],
    "reviews": ["reviews", "review_count", "google_review", "google_reviews", "num_reviews", "review"],
    "notes": ["notes", "note", "remarks", "description", "comment", "comments"],
}

HIGH_VALUE = [
    "real estate", "property", "builder", "developer", "architect", "interior",
    "construction", "contractor", "clinic", "dental", "dentist", "doctor", "hospital",
    "diagnostic", "salon", "spa", "beauty", "gym", "fitness", "yoga", "restaurant",
    "cafe", "bakery", "hotel", "resort", "travel", "tour", "lawyer", "advocate",
    "legal", "chartered accountant", "consultant", "consultancy", "agency",
    "education", "coaching", "tuition", "school", "college", "institute",
    "e-commerce", "ecommerce", "retail", "boutique", "shop", "store", "manufacturer",
    "supplier", "distributor", "wholesale", "event", "wedding", "photographer",
    "catering", "logistics", "transport", "automobile", "car", "bike", "showroom",
    "jewellery", "jewelry", "fashion", "clothing", "electronics", "furniture",
    "hardware", "pharma", "medical", "veterinary", "pet", "dealer",
]

MEDIUM_VALUE = [
    "service", "professional", "business", "company", "firm", "office", "group",
    "enterprise", "industries", "trading", "solutions", "studio", "academy",
]

STRING_COLUMNS = [
    "lead_id", "name", "first_name", "email", "phone", "business_type", "website",
    "city", "notes", "priority", "channel", "reason",
]


def normalize_column(name):
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def clean_text(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    text = re.sub(r"\s+", " ", str(value)).strip()
    if text.lower() in ("nan", "none", "null"):
        return ""
    return text


def _read_text(text):
    text = text.lstrip("\ufeff")
    if not text.strip():
        raise ValueError("The file is empty.")
    separator = None
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
        separator = dialect.delimiter
    except csv.Error:
        separator = None
    return pd.read_csv(io.StringIO(text), sep=separator, engine="python")


def _read_bytes(data, name):
    suffix = Path(name or "").suffix.lower()
    if suffix in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(io.BytesIO(data))
    if suffix == ".docx":
        return _read_docx(data)
    if suffix in (".csv", ".tsv", ".txt"):
        return _read_text(data.decode("utf-8-sig", errors="replace"))
    try:
        return pd.read_excel(io.BytesIO(data))
    except Exception:
        return _read_text(data.decode("utf-8-sig", errors="replace"))


def _dedupe_headers(headers):
    seen = {}
    result = []
    for index, header in enumerate(headers):
        label = clean_text(header) or f"Column {index + 1}"
        if label in seen:
            seen[label] += 1
            label = f"{label} {seen[label]}"
        else:
            seen[label] = 1
        result.append(label)
    return result


def _looks_like_header(row):
    generic = {
        "name", "full name", "business name", "company", "company name", "email", "e-mail", "mail",
        "email id", "phone", "phone number", "mobile", "mobile number", "contact", "contact number",
        "business type", "category", "industry", "city", "location", "area", "website", "rating",
        "reviews", "notes",
    }
    cells = [clean_text(cell).lower() for cell in row]
    if len(cells) == 1:
        return cells[0] in generic
    matches = sum(1 for cell in cells if cell in generic or any(
        word in cell for word in ("email", "mail", "phone", "mobile", "whatsapp")
    ))
    return matches >= 2


def _read_docx(data):
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("DOCX support needs python-docx. Run: pip install python-docx") from exc

    document = Document(io.BytesIO(data))
    frames = []
    for table in document.tables:
        rows = [[clean_text(cell.text) for cell in row.cells] for row in table.rows]
        rows = [row for row in rows if any(row)]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        if _looks_like_header(rows[0]):
            headers = _dedupe_headers(rows[0])
            data_rows = rows[1:]
        else:
            headers = [f"Column {index + 1}" for index in range(width)]
            data_rows = rows
        frames.append(pd.DataFrame(data_rows, columns=headers))
    if frames:
        return pd.concat(frames, ignore_index=True, sort=False)

    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    if not text.strip():
        raise ValueError("No tables or text found in the DOCX file.")
    return _read_text(text)


def load_data(source):
    if hasattr(source, "getvalue"):
        return _read_bytes(source.getvalue(), getattr(source, "name", ""))
    text = str(source).strip()
    if text.startswith(("http://", "https://")):
        url = text
        if "docs.google.com/spreadsheets" in url and "/export" not in url:
            base = url.split("/edit")[0] + "/export?format=csv"
            gid = re.search(r"[#&]gid=(\d+)", url)
            if gid:
                base += f"&gid={gid.group(1)}"
            url = base
        if url.lower().endswith((".xlsx", ".xls")) or "format=xlsx" in url:
            return pd.read_excel(url)
        return pd.read_csv(url)
    path = Path(text)
    if path.exists():
        if path.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
            return pd.read_excel(path)
        if path.suffix.lower() == ".docx":
            return _read_docx(path.read_bytes())
        return _read_text(path.read_text(encoding="utf-8-sig", errors="replace"))
    return _read_text(text)


def detect_mapping(df):
    normalized = {normalize_column(col): col for col in df.columns}
    used = set()
    mapping = {}

    for canonical, aliases in ALIASES.items():
        found = None
        for alias in aliases:
            key = normalize_column(alias)
            if key in normalized and normalized[key] not in used:
                found = normalized[key]
                break
        if found:
            mapping[canonical] = found
            used.add(found)

    fuzzy = {
        "email": ("email", "mail"),
        "phone": ("phone", "mobile", "whatsapp", "tel", "cell", "contact"),
        "name": ("company", "business", "firm", "agency", "name"),
        "business_type": ("industry", "category", "sector", "niche", "type"),
        "website": ("website", "site", "domain"),
        "city": ("city", "location", "area", "address"),
        "rating": ("rating", "stars"),
        "reviews": ("review",),
        "notes": ("note", "remark", "comment", "description"),
    }
    for canonical, hints in fuzzy.items():
        if canonical in mapping:
            continue
        for key, original in normalized.items():
            if original in used:
                continue
            if canonical == "name" and ("type" in key or "email" in key or "phone" in key):
                continue
            if any(hint in key for hint in hints):
                mapping[canonical] = original
                used.add(original)
                break

    return {canonical: mapping.get(canonical) for canonical in ALIASES}


def normalize_email(value):
    raw = clean_text(value)
    if not raw:
        return ""
    candidates = [part for part in re.split(r"[;,\s/<>()\[\]]+", raw) if "@" in part]
    for candidate in candidates:
        candidate = candidate.strip(" .,;:").lower()
        if not candidate:
            continue
        try:
            return validate_email(candidate, check_deliverability=False).normalized
        except EmailNotValidError:
            continue
    return ""


def normalize_phone(value, default_country="91"):
    raw = clean_text(value)
    if not raw:
        return ""
    if raw.lower().startswith(("n/a", "na", "none", "-")):
        return ""
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return ""
    if digits.startswith("00"):
        digits = digits[2:]
    country = re.sub(r"\D", "", str(default_country or ""))
    if country:
        if len(digits) == 10:
            digits = country + digits
        elif len(digits) == 11 and digits.startswith("0"):
            digits = country + digits[1:]
        elif len(digits) == 12 and digits.startswith("0" + country):
            digits = digits[1:]
    if len(digits) < 8 or len(digits) > 15:
        return ""
    return "+" + digits


def only_digits(value):
    return re.sub(r"\D", "", str(value or ""))


def parse_rating(value):
    text = clean_text(value)
    if not text:
        return None
    match = re.search(r"(\d(?:\.\d+)?)", text)
    if not match:
        return None
    try:
        rating = float(match.group(1))
    except ValueError:
        return None
    if 0 < rating <= 5:
        return round(rating, 1)
    return None


def parse_reviews(value):
    text = clean_text(value)
    if not text:
        return None
    match = re.search(r"reviews?\s*[:\-]?\s*([\d,]+)", text, flags=re.IGNORECASE)
    if match:
        number = match.group(1)
    else:
        numbers = re.findall(r"[\d,]+", text)
        if not numbers:
            return None
        number = numbers[-1]
    try:
        return int(number.replace(",", ""))
    except ValueError:
        return None


def prettify_name(value):
    name = re.sub(r"\s*\([^)]*\)", " ", clean_text(value)).strip()
    name = re.sub(r"\s+", " ", name)
    if name and name == name.upper():
        name = name.title()
    return name


def lead_id_for(name, email, phone):
    base = "|".join([email.lower(), only_digits(phone), name.lower()])
    return hashlib.sha1(base.encode("utf-8")).hexdigest()[:16]


def _value(row, mapping, key):
    column = mapping.get(key)
    if not column:
        return ""
    try:
        if column not in row.index:
            return ""
        return row[column]
    except Exception:
        return ""


def score_lead(lead):
    score = 15
    reasons = []

    email = clean_text(lead.get("email"))
    phone = clean_text(lead.get("phone"))
    if email:
        score += 20
        reasons.append("email available")
    if phone:
        score += 15
        reasons.append("phone available")

    business_type = clean_text(lead.get("business_type")).lower()
    if business_type:
        if any(word in business_type for word in HIGH_VALUE):
            score += 25
            reasons.append("high-fit industry")
        elif any(word in business_type for word in MEDIUM_VALUE):
            score += 15
            reasons.append("service business")
        else:
            score += 8
    else:
        score += 5

    rating = lead.get("rating")
    reviews = lead.get("reviews")
    try:
        rating = float(rating) if rating not in (None, "") else None
    except (TypeError, ValueError):
        rating = None
    try:
        reviews = int(float(reviews)) if reviews not in (None, "") else None
    except (TypeError, ValueError):
        reviews = None

    if reviews:
        if reviews >= 100:
            score += 10
            reasons.append("strong review volume")
        elif reviews >= 30:
            score += 7
        elif reviews >= 5:
            score += 4
        else:
            score += 2
    if rating:
        if rating >= 4.5:
            score += 10
            reasons.append("high rating")
        elif rating >= 4:
            score += 6
        else:
            score += 3

    if clean_text(lead.get("name")):
        score += 5
    website = clean_text(lead.get("website"))
    if website:
        score += 8
        reasons.append("website can be improved")
    else:
        score += 5

    score = max(0, min(100, int(round(score))))
    if score >= 75:
        priority = "High"
    elif score >= 55:
        priority = "Medium"
    else:
        priority = "Low"

    if email and phone:
        channel = "Email + WhatsApp"
    elif email:
        channel = "Email"
    elif phone:
        channel = "WhatsApp"
    else:
        channel = "Needs contact"

    reason = " - ".join(reasons) if reasons else "limited data"
    return score, priority, channel, reason


def process_leads(df, mapping, default_country="91"):
    df = df.dropna(how="all")
    records = []
    seen_emails = set()
    seen_phones = set()
    dropped = 0
    duplicates = 0

    for _, row in df.iterrows():
        raw_name = _value(row, mapping, "name")
        name = prettify_name(raw_name)
        email = normalize_email(_value(row, mapping, "email"))
        phone = normalize_phone(_value(row, mapping, "phone"), default_country)
        business_type = clean_text(_value(row, mapping, "business_type"))
        website = clean_text(_value(row, mapping, "website"))
        city = clean_text(_value(row, mapping, "city"))
        notes = clean_text(_value(row, mapping, "notes"))
        first_name = clean_text(_value(row, mapping, "first_name"))

        if not (name or email or phone):
            dropped += 1
            continue
        if email and email in seen_emails:
            duplicates += 1
            continue
        if phone and phone in seen_phones:
            duplicates += 1
            continue

        rating_raw = clean_text(_value(row, mapping, "rating"))
        reviews_raw = clean_text(_value(row, mapping, "reviews"))
        rating = parse_rating(rating_raw) if rating_raw else None
        reviews = parse_reviews(reviews_raw) if reviews_raw else None
        if rating is None and reviews_raw and "|" in reviews_raw:
            rating = parse_rating(reviews_raw)

        if not first_name and name:
            first_name = name.split()[0]

        lead = {
            "lead_id": lead_id_for(name, email, phone),
            "selected": True,
            "name": name or (first_name or "Unknown"),
            "first_name": first_name,
            "email": email,
            "phone": phone,
            "business_type": business_type,
            "website": website,
            "city": city,
            "rating": rating,
            "reviews": reviews,
            "notes": notes,
        }
        score, priority, channel, reason = score_lead(lead)
        lead.update({"score": score, "priority": priority, "channel": channel, "reason": reason})
        records.append(lead)

        if email:
            seen_emails.add(email)
        if phone:
            seen_phones.add(phone)

    result = pd.DataFrame(records)
    if result.empty:
        result = pd.DataFrame(columns=[
            "lead_id", "selected", "name", "first_name", "email", "phone", "business_type",
            "website", "city", "rating", "reviews", "notes", "score", "priority", "channel", "reason",
        ])
    result = ensure_schema(result)
    stats = {
        "rows_read": int(len(df)),
        "kept": int(len(result)),
        "dropped": int(dropped),
        "duplicates": int(duplicates),
    }
    return result, stats


def ensure_schema(df):
    df = df.copy()
    for column in STRING_COLUMNS:
        if column not in df.columns:
            df[column] = ""
        df[column] = df[column].fillna("").astype(str).replace("nan", "")
    if "selected" not in df.columns:
        df["selected"] = True
    df["selected"] = df["selected"].fillna(True).astype(bool)
    if "score" not in df.columns:
        df["score"] = 0
    df["score"] = pd.to_numeric(df["score"], errors="coerce").fillna(0).astype(int)
    for column in ("rating", "reviews"):
        if column not in df.columns:
            df[column] = None
        df[column] = pd.to_numeric(df[column], errors="coerce")
    return df


def recalculate_scores(df):
    df = df.copy()
    if df.empty:
        return ensure_schema(df)
    values = []
    for _, row in df.iterrows():
        values.append(score_lead(row.to_dict()))
    df["score"] = [item[0] for item in values]
    df["priority"] = [item[1] for item in values]
    df["channel"] = [item[2] for item in values]
    df["reason"] = [item[3] for item in values]
    return ensure_schema(df)
